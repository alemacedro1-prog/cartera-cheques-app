from __future__ import annotations

import io
import re
import unicodedata
from datetime import date

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.portfolio import (
    ALLOWED_TYPES,
    PENDING_COLLECTION_STATES,
    _clean,
    _date_value,
    _extract_observation_receipt,
    _first,
    _identifier,
    build_portfolio,
    instrument_amount,
    instrument_state,
)
from utils.consolidation import SOURCE_ROWS, cheque_number, rejected_then_accredited, select_source_rows


MOVEMENT_LINK_STATES = ("Con recibo", "Sin recibo", "A revisar")
PROCESSING_VERSION = "consolidation-v10"
MANUAL_ORIGIN = "MANU - acreditado fuera del concentrador"
CONTROL_RECORD_TYPES = ("Cheque", "Movimiento bancario")
RECONCILIATION_STATES = (
    "Coincide recibo e importe",
    "Mismo recibo, importe diferente",
    "Solo en cheques",
    "Solo en movimientos bancarios",
)


def _normalized_text(value) -> str:
    text = unicodedata.normalize("NFKD", _clean(value))
    return "".join(character for character in text if not unicodedata.combining(character)).upper()


def payment_method_group(value) -> str:
    """Agrupa medios no cheque sin perder la denominación original."""
    original = _clean(value).upper()
    normalized = _normalized_text(value)
    if not normalized:
        return ""
    if normalized in {"EF", "EFE", "CASH"} or "EFECT" in normalized:
        return "Efectivo"
    if normalized in {"TR", "TRF", "TRANSF", "TRAN"} or "TRANSFER" in normalized:
        return "Transferencia"
    if normalized in {"DEP", "DEPO", "DB"} or "DEPOSIT" in normalized:
        return "Depósito"
    return original or "Otro"


def _receipt_key(value: str) -> str:
    """Compara números equivalentes aunque cambien espacios o separadores."""
    return re.sub(r"[^A-Z0-9]", "", _clean(value).upper())


def _typed_related_receipt(row: pd.Series) -> str:
    related_number = _identifier(row.get("Nro Cpb Relacionado"))
    related_type = _normalized_text(_first(row, "Tipo Cpb Relacionado", "MCR-Tipo Cpb Relacionado"))
    return related_number if related_number and ("REC" in related_type or "COB" in related_type) else ""


def resolve_movement_receipt(row: pd.Series, method_group: str) -> dict[str, str]:
    """Resuelve recibos de medios no cheque y expone incompatibilidades.

    Las prioridades varían por medio. A diferencia de los cheques, el campo
    MCR-Número de recibo se considera una señal válida, pero nunca oculta una
    contradicción con las fuentes de cobranza.
    """
    candidates = {
        "Observación": _extract_observation_receipt(_clean(row.get("Observación"))),
        "Nro Cpb Relación": _identifier(row.get("Nro Cpb Relación")),
        "Comprobante relacionado tipificado": _typed_related_receipt(row),
        "MCR-Número de recibo": _identifier(row.get("MCR-Número de recibo")),
    }
    candidates = {source: value for source, value in candidates.items() if value}

    priorities = {
        "Efectivo": (
            "MCR-Número de recibo",
            "Nro Cpb Relación",
            "Observación",
            "Comprobante relacionado tipificado",
        ),
        "Transferencia": (
            "Nro Cpb Relación",
            "Observación",
            "Comprobante relacionado tipificado",
            "MCR-Número de recibo",
        ),
        "Depósito": (
            "Nro Cpb Relación",
            "Observación",
            "MCR-Número de recibo",
            "Comprobante relacionado tipificado",
        ),
    }
    priority = priorities.get(
        method_group,
        ("Nro Cpb Relación", "Observación", "Comprobante relacionado tipificado", "MCR-Número de recibo"),
    )
    if not candidates:
        return {
            "Estado vínculo": "Sin recibo",
            "Recibo relacionado": "",
            "Fuente del vínculo": "Sin vínculo detectado",
            "Fuentes detectadas": "",
        }

    distinct = {}
    for source, value in candidates.items():
        distinct.setdefault(_receipt_key(value), []).append(source)
    chosen_source = next(source for source in priority if source in candidates)
    chosen = candidates[chosen_source]
    detected = " · ".join(f"{source}: {value}" for source, value in candidates.items())
    if len(distinct) > 1:
        return {
            "Estado vínculo": "A revisar",
            "Recibo relacionado": chosen,
            "Fuente del vínculo": "Fuentes incompatibles",
            "Fuentes detectadas": detected,
        }
    matching_sources = " + ".join(candidates)
    return {
        "Estado vínculo": "Con recibo",
        "Recibo relacionado": chosen,
        "Fuente del vínculo": matching_sources,
        "Fuentes detectadas": detected,
    }


def _cheque_control_records(raw: pd.DataFrame, portfolio: pd.DataFrame) -> list[dict]:
    """Adapta la cartera al mismo esquema del control de movimientos."""
    if portfolio.empty:
        return []
    raw_by_row = raw.set_index("Fila fuente", drop=False) if "Fila fuente" in raw else pd.DataFrame()
    records: list[dict] = []
    for _, cheque in portfolio.iterrows():
        source_row = int(cheque["Fila fuente"])
        source = raw_by_row.loc[source_row] if not raw_by_row.empty and source_row in raw_by_row.index else pd.Series(dtype=object)
        if isinstance(source, pd.DataFrame):
            source = source.iloc[0]
        receipt = _identifier(cheque.get("Recibo relacionado"))
        receipt_source = _clean(cheque.get("Fuente del vínculo")) or "Sin vínculo detectado"
        numeric_amount = pd.to_numeric(cheque.get("Importe"), errors="coerce")
        records.append(
            {
                "ID movimiento": _clean(cheque.get("ID cartera")),
                "Fila fuente": source_row,
                SOURCE_ROWS: cheque.get(SOURCE_ROWS, str(source_row)),
                "Cantidad registros origen": int(cheque.get("Cantidad registros origen", 1)),
                "Duplicados consolidados": int(cheque.get("Duplicados consolidados", 0)),
                "Historial de estados": _clean(cheque.get("Historial de estados")),
                "Rechazado luego acreditado": bool(cheque.get("Rechazado luego acreditado", False)),
                "Tipo de registro": "Cheque",
                "Fecha": _date_value(cheque.get("Fecha prevista de cobro")) if pd.notna(cheque.get("Fecha prevista de cobro")) else _date_value(cheque.get("Fecha ingreso / pago")),
                "Fecha ingreso / pago": _date_value(cheque.get("Fecha ingreso / pago")),
                "Fecha acreditación": _date_value(cheque.get("Fecha acreditación")),
                "Fecha vencimiento": _date_value(cheque.get("Fecha vencimiento")),
                "Cliente": _clean(cheque.get("Cliente")),
                "CUIT cliente": _identifier(cheque.get("CUIT cliente")),
                "Importe": float(numeric_amount) if not pd.isna(numeric_amount) else 0.0,
                "Medio de pago": _clean(cheque.get("Tipo")),
                "Grupo medio de pago": "Cheque",
                "Banco / cuenta": _clean(cheque.get("Banco depósito")) or _identifier(cheque.get("Banco cheque")),
                "Subcuenta": _identifier(_first(source, "SubCuenta banco", "MCR-Subcuenta depósito")),
                "Estado vínculo": "Con recibo" if _clean(cheque.get("Estado recibo")) == "Tomado" else "Sin recibo",
                "Recibo relacionado": receipt,
                "Fuente del vínculo": receipt_source,
                "Fuentes detectadas": f"{receipt_source}: {receipt}" if receipt else "",
                "Estado operativo": _clean(cheque.get("Estado calculado")),
                "Estado original": _clean(cheque.get("Estado original")),
                "Código estado": _clean(cheque.get("Código estado")),
                "Origen": "CONRENPF - cheque consolidado",
                "MANU acreditado fuera": False,
                "N° cheque / eCheq": cheque_number(cheque.get("N° cheque / eCheq")),
                "Nro Cpb Relación": _identifier(cheque.get("Nro Cpb Relación")),
                "Nro Cpb Relacionado": _identifier(source.get("Nro Cpb Relacionado")),
                "MCR-Número de recibo": _identifier(source.get("MCR-Número de recibo")),
                "Observación": _clean(cheque.get("Observaciones")),
                "N° operación": _identifier(cheque.get("N° operación")),
                "MCR-ID pago": _identifier(cheque.get("MCR-ID pago")),
                "MCR-ID instrumento": _identifier(cheque.get("MCR-ID instrumento")),
            }
        )
    return records


def build_movement_control(raw: pd.DataFrame, portfolio: pd.DataFrame | None = None) -> pd.DataFrame:
    """Construye un control unificado de cheques y movimientos bancarios."""
    if "MCR-Medio de pago" not in raw:
        return pd.DataFrame()
    if portfolio is None:
        portfolio = build_portfolio(raw)
    media = raw["MCR-Medio de pago"].map(lambda value: _clean(value).upper())
    target = raw[media.ne("") & ~media.isin(ALLOWED_TYPES)].copy()
    if "Fila fuente" not in target:
        target["Fila fuente"] = target.index + 2

    records: list[dict] = _cheque_control_records(raw, portfolio)
    for number, (_, row) in enumerate(target.iterrows(), start=1):
        raw_method = _clean(row.get("MCR-Medio de pago")).upper()
        method_group = payment_method_group(raw_method)
        link = resolve_movement_receipt(row, method_group)
        amount = instrument_amount(row)
        movement_date = _date_value(
            _first(
                row,
                "MCR-Fecha pago",
                "Fecha Movimiento",
                "Fecha Depósito",
                "MCR-Fecha rendición",
                "Fecha Rendición",
            )
        )
        original_state, state_code = instrument_state(row)
        manual_external = raw_method == "MANU" and state_code not in {"AC", "RC", "PS"}
        operational_state = {
            "AC": "Acreditado",
            "RC": "Rechazado",
            "PS": "Pendiente de acreditación",
            "RE": "Rescatado",
            "PE": "Pendiente",
            "P": "Pendiente",
        }.get(state_code, "Sin estado informado")
        if manual_external:
            operational_state = "Acreditado"
        records.append(
            {
                "ID movimiento": f"MOV-{number:04d}",
                "Fila fuente": int(row["Fila fuente"]),
                SOURCE_ROWS: str(int(row["Fila fuente"])),
                "Cantidad registros origen": 1,
                "Duplicados consolidados": 0,
                "Historial de estados": f"Fila {int(row['Fila fuente'])}: {original_state or 'sin estado'}",
                "Rechazado luego acreditado": False,
                "Tipo de registro": "Movimiento bancario",
                "Fecha": movement_date,
                "Fecha ingreso / pago": movement_date,
                "Fecha acreditación": pd.NaT,
                "Fecha vencimiento": pd.NaT,
                "Cliente": _clean(_first(row, "MCR-Nombre cliente", "Nombre")),
                "CUIT cliente": _identifier(row.get("MCR-CUIT cliente")),
                "Importe": float(amount) if not pd.isna(amount) else 0.0,
                "Medio de pago": raw_method,
                "Grupo medio de pago": method_group,
                "Banco / cuenta": _clean(_first(row, "Nombre del banco", "MCR-Banco depósito", "MCR-Banco")),
                "Subcuenta": _identifier(_first(row, "SubCuenta banco", "MCR-Subcuenta depósito")),
                "Estado vínculo": link["Estado vínculo"],
                "Recibo relacionado": link["Recibo relacionado"],
                "Fuente del vínculo": link["Fuente del vínculo"],
                "Fuentes detectadas": link["Fuentes detectadas"],
                "Estado operativo": operational_state,
                "Estado original": original_state,
                "Código estado": state_code,
                "Origen": MANUAL_ORIGIN if manual_external else "CONRENPF - movimiento bancario",
                "MANU acreditado fuera": manual_external,
                "N° cheque / eCheq": next((cheque_number(row.get(column)) for column in ("MCR-Número de cheque", "Nro de Cheque", "N° cheque / eCheq") if cheque_number(row.get(column))), ""),
                "Nro Cpb Relación": _identifier(row.get("Nro Cpb Relación")),
                "Nro Cpb Relacionado": _identifier(row.get("Nro Cpb Relacionado")),
                "MCR-Número de recibo": _identifier(row.get("MCR-Número de recibo")),
                "Observación": _clean(row.get("Observación")),
                "N° operación": _identifier(_first(row, "MCR-Número operación", "MCR-Nro instrumento", "Nro de Movimiento")),
                "MCR-ID pago": _identifier(row.get("MCR-ID pago")),
                "MCR-ID instrumento": _identifier(row.get("MCR-ID instrumento")),
            }
        )
    if not records:
        return pd.DataFrame()
    return pd.DataFrame.from_records(records).sort_values("Fila fuente").reset_index(drop=True)


def pending_movements_detail(movements: pd.DataFrame) -> pd.DataFrame:
    columns = ["Cliente", "Método de pago", "Fecha", "Importe", "N° cheque / eCheq", "Recibo relacionado", "Estado", "Origen", SOURCE_ROWS]
    if movements.empty:
        return pd.DataFrame(columns=columns)
    pending = movements[movements["Estado operativo"].isin(PENDING_COLLECTION_STATES)].copy()
    return pending.rename(columns={"Medio de pago": "Método de pago", "Estado operativo": "Estado"}).reindex(columns=columns).sort_values("Fecha", na_position="last").reset_index(drop=True)


def manual_accredited_movements(movements: pd.DataFrame) -> pd.DataFrame:
    mask = movements.get("MANU acreditado fuera", pd.Series(False, index=movements.index)).fillna(False).astype(bool)
    return movements[mask].copy()


def operational_export_tables(movements: pd.DataFrame, raw: pd.DataFrame, cutoff: date) -> dict[str, pd.DataFrame]:
    """Mismas categorías y mismo alcance que la selección, con todo su historial."""
    cheque_mask = movements.get("Tipo de registro", pd.Series(index=movements.index, dtype=object)).eq("Cheque")
    cheque_raw = select_source_rows(movements[cheque_mask], raw)
    portfolio = build_portfolio(cheque_raw, cutoff)
    if portfolio.empty:
        portfolio = pd.DataFrame(columns=["Cliente", "N° cheque / eCheq", "Importe", "Estado calculado", "Estado final", SOURCE_ROWS])
    states = portfolio.get("Estado calculado", pd.Series(index=portfolio.index, dtype=object))
    return {
        "Cartera consolidada": portfolio,
        "Movimientos pendientes": pending_movements_detail(movements),
        "Rechazados efectivos": portfolio[states.eq("Rechazado")].copy(),
        "Rechazados luego acreditados": rejected_then_accredited(portfolio),
        "MANU acreditados": manual_accredited_movements(movements),
    }


def movement_link_summary(movements: pd.DataFrame) -> pd.DataFrame:
    columns = ["Estado vínculo", "Cantidad", "Importe", "Participación cantidad", "Participación importe"]
    if movements.empty:
        return pd.DataFrame(columns=columns)
    work = movements.copy()
    work["Importe"] = pd.to_numeric(work["Importe"], errors="coerce").fillna(0)
    summary = work.groupby("Estado vínculo", as_index=False).agg(Cantidad=("Importe", "size"), Importe=("Importe", "sum"))
    summary = summary.set_index("Estado vínculo").reindex(MOVEMENT_LINK_STATES, fill_value=0).reset_index()
    total_count = int(summary["Cantidad"].sum())
    total_amount = float(summary["Importe"].sum())
    summary["Participación cantidad"] = summary["Cantidad"] / total_count if total_count else 0.0
    summary["Participación importe"] = summary["Importe"] / total_amount if total_amount else 0.0
    return summary[columns]


def movement_type_summary(movements: pd.DataFrame) -> pd.DataFrame:
    columns = ["Tipo de registro", "Cantidad", "Importe", "Participación importe"]
    if movements.empty:
        return pd.DataFrame(columns=columns)
    work = movements.copy()
    work["Importe"] = pd.to_numeric(work["Importe"], errors="coerce").fillna(0)
    summary = work.groupby("Tipo de registro", as_index=False).agg(
        Cantidad=("Importe", "size"), Importe=("Importe", "sum")
    )
    summary = summary.set_index("Tipo de registro").reindex(CONTROL_RECORD_TYPES, fill_value=0).reset_index()
    total_amount = float(summary["Importe"].sum())
    summary["Participación importe"] = summary["Importe"] / total_amount if total_amount else 0.0
    return summary[columns]


def receipt_reconciliation(movements: pd.DataFrame) -> pd.DataFrame:
    """Cruza cheques y movimientos bancarios por un recibo no conflictivo."""
    columns = [
        "Recibo",
        "Cheques",
        "Importe cheques",
        "Movimientos bancarios",
        "Importe movimientos bancarios",
        "Diferencia",
        "Resultado",
    ]
    if movements.empty:
        return pd.DataFrame(columns=columns)
    work = movements.copy()
    receipts = work.get("Recibo relacionado", pd.Series(index=work.index, dtype=object)).map(_identifier)
    states = work.get("Estado vínculo", pd.Series(index=work.index, dtype=object))
    work = work[receipts.ne("") & states.eq("Con recibo")].copy()
    if work.empty:
        return pd.DataFrame(columns=columns)
    work["Recibo"] = receipts.loc[work.index]
    work["Clave recibo"] = work["Recibo"].map(_receipt_key)
    work["Importe"] = pd.to_numeric(work["Importe"], errors="coerce").fillna(0)

    records: list[dict] = []
    for _, receipt_rows in work.groupby("Clave recibo", sort=True):
        cheque_rows = receipt_rows[receipt_rows["Tipo de registro"].eq("Cheque")]
        bank_rows = receipt_rows[receipt_rows["Tipo de registro"].eq("Movimiento bancario")]
        cheque_amount = float(cheque_rows["Importe"].sum())
        bank_amount = float(bank_rows["Importe"].sum())
        difference = bank_amount - cheque_amount
        if not cheque_rows.empty and not bank_rows.empty:
            result = "Coincide recibo e importe" if abs(difference) <= 0.01 else "Mismo recibo, importe diferente"
        elif not cheque_rows.empty:
            result = "Solo en cheques"
        else:
            result = "Solo en movimientos bancarios"
        records.append(
            {
                "Recibo": receipt_rows.iloc[0]["Recibo"],
                "Cheques": len(cheque_rows),
                "Importe cheques": cheque_amount,
                "Movimientos bancarios": len(bank_rows),
                "Importe movimientos bancarios": bank_amount,
                "Diferencia": difference,
                "Resultado": result,
            }
        )
    return pd.DataFrame.from_records(records, columns=columns).sort_values(
        ["Resultado", "Diferencia", "Recibo"], ascending=[True, False, True]
    ).reset_index(drop=True)


def _excel_safe(frame: pd.DataFrame) -> pd.DataFrame:
    safe = frame.copy()
    for column in safe.select_dtypes(include=["object", "string"]).columns:
        safe[column] = safe[column].map(
            lambda value: "'" + value
            if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@"))
            else value
        )
    return safe


def export_movements_excel(movements: pd.DataFrame, raw: pd.DataFrame, analysis_date: date) -> bytes:
    scoped_raw = select_source_rows(movements, raw)
    operational_tables = operational_export_tables(movements, scoped_raw, analysis_date)
    amounts = pd.to_numeric(movements.get("Importe", pd.Series(dtype=float)), errors="coerce").fillna(0)
    states = movements.get("Estado vínculo", pd.Series(index=movements.index, dtype=object))
    reconciliation = receipt_reconciliation(movements)
    summary_rows = [
        ("Fecha de análisis", pd.Timestamp(analysis_date)),
        ("Registros controlados", len(movements)),
        ("Importe total", amounts.sum()),
        ("Duplicados de cheques consolidados", int(movements.get("Duplicados consolidados", pd.Series(dtype=int)).sum())),
        ("Movimientos pendientes", len(operational_tables["Movimientos pendientes"])),
        ("Rechazados efectivos", len(operational_tables["Rechazados efectivos"])),
        ("Rechazados luego acreditados", len(operational_tables["Rechazados luego acreditados"])),
        ("MANU acreditados fuera del concentrador", len(operational_tables["MANU acreditados"])),
    ]
    record_types = movements.get("Tipo de registro", pd.Series(index=movements.index, dtype=object))
    for record_type in CONTROL_RECORD_TYPES:
        mask = record_types.eq(record_type)
        summary_rows.extend(
            [
                (f"{record_type} - cantidad", int(mask.sum())),
                (f"{record_type} - importe", amounts[mask].sum()),
            ]
        )
    for state in MOVEMENT_LINK_STATES:
        mask = states.eq(state)
        summary_rows.extend(
            [
                (f"{state} - cantidad", int(mask.sum())),
                (f"{state} - importe", amounts[mask].sum()),
            ]
        )
    for state in RECONCILIATION_STATES:
        summary_rows.append((f"Cruce - {state}", int(reconciliation["Resultado"].eq(state).sum())))
    summary = pd.DataFrame(summary_rows, columns=["Indicador", "Valor"])

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Resumen", index=False)
        _excel_safe(movements).to_excel(writer, sheet_name="Movimientos", index=False)
        _excel_safe(reconciliation).to_excel(writer, sheet_name="Cruce por recibo", index=False)
        for sheet_name, table in operational_tables.items():
            _excel_safe(table).to_excel(writer, sheet_name=sheet_name, index=False)
        _excel_safe(scoped_raw).to_excel(writer, sheet_name="Datos fuente filtrados", index=False)
        for sheet_name in writer.book.sheetnames:
            ws = writer.book[sheet_name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="17365D")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for cells in ws.columns:
                letter = get_column_letter(cells[0].column)
                sample = [str(cell.value or "") for cell in list(cells)[:150]]
                ws.column_dimensions[letter].width = min(max(max(map(len, sample), default=8) + 2, 11), 38)
                header = str(cells[0].value or "")
                if header.startswith("Fecha"):
                    for cell in cells[1:]:
                        cell.number_format = "dd/mm/yyyy"
                elif header.startswith("Importe") and header != "Importe informado":
                    for cell in cells[1:]:
                        cell.number_format = '$#,##0.00;[Red]($#,##0.00);-'
        summary_ws = writer.book["Resumen"]
        for row_number, indicator in enumerate(summary["Indicador"], start=2):
            value_cell = summary_ws.cell(row=row_number, column=2)
            if indicator == "Fecha de análisis":
                value_cell.number_format = "dd/mm/yyyy"
            elif indicator == "Importe total" or indicator.endswith(" - importe"):
                value_cell.number_format = '$#,##0.00;[Red]($#,##0.00);-'
            else:
                value_cell.number_format = "#,##0"
        movement_ws = writer.book["Movimientos"]
        if "Importe" in movements:
            amount_letter = get_column_letter(list(movements.columns).index("Importe") + 1)
            for cell in movement_ws[amount_letter][1:]:
                cell.number_format = '$#,##0.00;[Red]($#,##0.00);-'
        for date_column in ("Fecha", "Fecha ingreso / pago", "Fecha acreditación", "Fecha vencimiento"):
            if date_column in movements:
                date_letter = get_column_letter(list(movements.columns).index(date_column) + 1)
                for cell in movement_ws[date_letter][1:]:
                    cell.number_format = "dd/mm/yyyy"
        reconciliation_ws = writer.book["Cruce por recibo"]
        for amount_column in ("Importe cheques", "Importe movimientos bancarios", "Diferencia"):
            if amount_column in reconciliation:
                amount_letter = get_column_letter(list(reconciliation.columns).index(amount_column) + 1)
                for cell in reconciliation_ws[amount_letter][1:]:
                    cell.number_format = '$#,##0.00;[Red]($#,##0.00);-'
    return output.getvalue()
